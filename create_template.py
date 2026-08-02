from docx import Document

def create_master_template():
    doc = Document()
    doc.add_heading('Master Template', 0)
    doc.add_paragraph('This is a basic template for the AI Book Writer.')
    doc.save('master_template.docx')
    print("Template created.")

if __name__ == "__main__":
    create_master_template()
