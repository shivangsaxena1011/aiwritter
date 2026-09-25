import os
import re
import hashlib
import logging
from typing import Dict, Any, List, Optional, Tuple, Set
from dataclasses import dataclass, field
import docx
from docx import Document

from backend.app.services.document.artifact_integrity import CountManifest

logger = logging.getLogger(__name__)


@dataclass
class IndependentAuditResult:
    publication_ready: bool
    docx_path: str
    total_docx_paragraphs: int
    evaluated_prose_paragraphs: int
    exact_duplicate_paragraphs: int
    near_duplicate_paragraphs: int
    exact_duplicate_paragraph_rate: float
    near_duplicate_paragraph_rate: float
    total_headings: int
    consecutive_duplicate_headings: int
    generic_headings: int
    structural_template_families: List[str]
    total_tables: int
    exact_duplicate_tables: int
    near_duplicate_tables: int
    semantic_duplicate_tables: int
    math_rendering_errors_in_tables: int
    total_figures: int
    figure_caption_mismatches: int
    math_rendering_artifacts_count: int
    count_reconciliation_valid: bool
    blocking_reasons: List[str]
    summary_counts: Dict[str, int]
    reconciliation_details: Dict[str, Any]

    def to_dict(self) -> Dict[str, Any]:
        return {
            "publication_ready": self.publication_ready,
            "docx_path": self.docx_path,
            "total_docx_paragraphs": self.total_docx_paragraphs,
            "evaluated_prose_paragraphs": self.evaluated_prose_paragraphs,
            "exact_duplicate_paragraphs": self.exact_duplicate_paragraphs,
            "near_duplicate_paragraphs": self.near_duplicate_paragraphs,
            "exact_duplicate_paragraph_rate": round(self.exact_duplicate_paragraph_rate, 4),
            "near_duplicate_paragraph_rate": round(self.near_duplicate_paragraph_rate, 4),
            "total_headings": self.total_headings,
            "consecutive_duplicate_headings": self.consecutive_duplicate_headings,
            "generic_headings": self.generic_headings,
            "structural_template_families": self.structural_template_families,
            "total_tables": self.total_tables,
            "exact_duplicate_tables": self.exact_duplicate_tables,
            "near_duplicate_tables": self.near_duplicate_tables,
            "semantic_duplicate_tables": self.semantic_duplicate_tables,
            "math_rendering_errors_in_tables": self.math_rendering_errors_in_tables,
            "total_figures": self.total_figures,
            "figure_caption_mismatches": self.figure_caption_mismatches,
            "math_rendering_artifacts_count": self.math_rendering_artifacts_count,
            "count_reconciliation_valid": self.count_reconciliation_valid,
            "blocking_reasons": self.blocking_reasons,
            "summary_counts": self.summary_counts,
            "reconciliation_details": self.reconciliation_details
        }


class IndependentArtifactAuditor:
    """
    Content Orchestration 2.2 — Independent Artifact Truth Engine.

    Guiding Principles:
    1. THE RENDERED DOCX IS THE ONLY AUTHORITATIVE TRUTH.
    2. Independent inspection from scratch using python-docx and XML parsing.
    3. Evaluates ALL prose paragraphs in the DOCX (no sampling or partial subsets).
    4. Evaluates ALL tables for exact, near, and semantic duplicates, plus math rendering artifacts.
    5. Strict 3-way count reconciliation: planned == assembled == rendered_docx.
    6. Identifies structural template family repetitions across topics.
    7. Identifies figure-caption semantic mismatches and math rendering artifacts.
    """

    GENERIC_HEADING_PATTERNS = [
        re.compile(r"^section\s+\d+$", re.IGNORECASE),
        re.compile(r"^subtopic\s+\d+$", re.IGNORECASE),
        re.compile(r"^untitled\s+section$", re.IGNORECASE),
        re.compile(r"^placeholder.*$", re.IGNORECASE),
        re.compile(r"^topic\s+\d+$", re.IGNORECASE)
    ]

    UNRENDERED_MATH_PATTERNS = [
        re.compile(r"\\[a-zA-Z]{2,}"),  # e.g. \frac, \sqrt, \text, \lambda
        re.compile(r"\$\$"),  # Display math delimiter unrendered
        re.compile(r"(?<!\\)\$[^\$\n]{2,}\$")  # Raw LaTeX inline math $...$ unrendered in table/text
    ]

    def __init__(self, near_threshold: float = 0.75, min_prose_words: int = 6):
        self.near_threshold = near_threshold
        self.min_prose_words = min_prose_words

    @staticmethod
    def _normalize(text: str) -> str:
        t = re.sub(r"[^\w\s]", "", text.lower())
        return " ".join(t.split())

    @staticmethod
    def _get_archetype(subtopic_heading: str) -> str:
        s = subtopic_heading.lower()
        if "concept" in s or "principle" in s or "motivation" in s or "foundation" in s:
            return "CONCEPT"
        if "formulation" in s or "governing" in s or "equation" in s:
            return "FORMULATION"
        if "derivation" in s or "boundary condition" in s or "proof" in s:
            return "DERIVATION"
        if "interpretation" in s or "eigenstate" in s or "amplitude" in s:
            return "INTERPRETATION"
        if "application" in s or "technology" in s or "relevance" in s or "engineering" in s:
            return "APPLICATION"
        if "limitation" in s or "assumption" in s or "scope" in s:
            return "LIMITATION"
        if "experiment" in s or "apparatus" in s or "diffraction" in s:
            return "EXPERIMENT"
        if "numerical" in s or "problem" in s or "exemplar" in s:
            return "NUMERICAL"
        return "GENERAL"

    def audit(
        self,
        docx_path: str,
        planned_manifest: Optional[CountManifest] = None,
        assembled_manifest: Optional[CountManifest] = None
    ) -> IndependentAuditResult:
        blocking_reasons: List[str] = []

        if not os.path.exists(docx_path):
            return IndependentAuditResult(
                publication_ready=False,
                docx_path=docx_path,
                total_docx_paragraphs=0,
                evaluated_prose_paragraphs=0,
                exact_duplicate_paragraphs=0,
                near_duplicate_paragraphs=0,
                exact_duplicate_paragraph_rate=0.0,
                near_duplicate_paragraph_rate=0.0,
                total_headings=0,
                consecutive_duplicate_headings=0,
                generic_headings=0,
                structural_template_families=[],
                total_tables=0,
                exact_duplicate_tables=0,
                near_duplicate_tables=0,
                semantic_duplicate_tables=0,
                math_rendering_errors_in_tables=0,
                total_figures=0,
                figure_caption_mismatches=0,
                math_rendering_artifacts_count=0,
                count_reconciliation_valid=False,
                blocking_reasons=[f"Missing file: {docx_path}"],
                summary_counts={},
                reconciliation_details={}
            )

        doc = Document(docx_path)

        # -------------------------------------------------------------
        # 1. PARAGRAPH & INVENTORY AUDIT
        # -------------------------------------------------------------
        total_docx_paragraphs = len(doc.paragraphs)
        headings: List[Dict[str, Any]] = []
        prose_paragraphs: List[Dict[str, Any]] = []
        figure_captions: List[Dict[str, Any]] = []

        chapter_headings: List[str] = []
        topic_headings: List[str] = []
        subtopic_headings: List[str] = []

        topic_structure_map: Dict[str, List[str]] = {}

        total_equations = 0
        total_figures = 0
        total_words = 0
        math_artifacts_in_text = 0

        current_chapter = "Front Matter"
        current_topic = ""
        current_subtopic = ""

        FRONT_BACK_H1 = {"preface", "table of contents", "appendix academic quality verification scorecard", "references"}

        for p_idx, p in enumerate(doc.paragraphs):
            p_xml = p._p.xml
            text = p.text.strip()
            style_name = p.style.name if p.style else ""
            words = len(text.split())
            total_words += words

            # Equation check
            has_eq = ("<m:oMath" in p_xml or "Cambria Math" in p_xml or "<m:f>" in p_xml)
            if has_eq:
                total_equations += 1

            # Drawing/Figure check
            has_drawing = ("<a:blip" in p_xml or "<w:drawing" in p_xml)
            if has_drawing:
                total_figures += 1

            is_heading = style_name.startswith("Heading")
            is_bullet = style_name.startswith("List")

            # Scan for unrendered math artifacts in prose text
            if not is_heading and not is_bullet and not has_eq:
                for pat in self.UNRENDERED_MATH_PATTERNS:
                    if pat.search(text):
                        math_artifacts_in_text += 1
                        break

            if is_heading:
                lvl = 1
                try:
                    lvl = int(style_name.replace("Heading", "").strip())
                except ValueError:
                    pass

                norm = self._normalize(text)
                headings.append({
                    "text": text,
                    "level": lvl,
                    "paragraph_index": p_idx,
                    "normalized": norm
                })

                if lvl == 1:
                    current_chapter = text
                    if norm not in FRONT_BACK_H1 and "appendix" not in norm:
                        chapter_headings.append(text)
                elif lvl == 2:
                    current_topic = text
                    if norm not in ("chapter overview objectives", "chapter overview and objectives", "academic bibliography"):
                        topic_headings.append(text)
                        if text not in topic_structure_map:
                            topic_structure_map[text] = []
                elif lvl == 3:
                    current_subtopic = text
                    if norm != "references":
                        subtopic_headings.append(text)
                        if current_topic and current_topic in topic_structure_map:
                            archetype = self._get_archetype(text)
                            topic_structure_map[current_topic].append(archetype)
            else:
                if text:
                    loc = f"{current_chapter} > {current_topic} > {current_subtopic}".strip(" >")

                    if text.startswith("Figure ") or text.startswith("Fig."):
                        figure_captions.append({
                            "text": text,
                            "topic": current_topic,
                            "subtopic": current_subtopic,
                            "location": loc
                        })
                    elif not is_bullet and not has_eq and words >= self.min_prose_words:
                        prose_paragraphs.append({
                            "index": p_idx,
                            "text": text,
                            "location": loc,
                            "words": words,
                            "normalized": self._normalize(text)
                        })

        # -------------------------------------------------------------
        # 2. HEADING INTEGRITY & CONSECUTIVE DUPLICATE AUDIT
        # -------------------------------------------------------------
        consecutive_duplicate_headings = 0
        generic_headings_count = 0
        duplicate_topic_headings = []
        seen_topics = set()

        for idx, h in enumerate(headings):
            if idx > 0:
                prev_h = headings[idx - 1]
                if h["normalized"] and h["normalized"] == prev_h["normalized"]:
                    consecutive_duplicate_headings += 1

            for pat in self.GENERIC_HEADING_PATTERNS:
                if pat.match(h["text"]):
                    generic_headings_count += 1
                    break

            if h["level"] == 2 and h["normalized"] and h["normalized"] != "chapter overview objectives":
                if h["normalized"] in seen_topics:
                    duplicate_topic_headings.append(h["text"])
                seen_topics.add(h["normalized"])

        if consecutive_duplicate_headings > 0:
            blocking_reasons.append(f"Found {consecutive_duplicate_headings} consecutive duplicate heading(s) in DOCX.")

        if generic_headings_count > 0:
            blocking_reasons.append(f"Found {generic_headings_count} generic fallback heading(s) in DOCX.")

        if len(duplicate_topic_headings) > 0:
            blocking_reasons.append(f"Found duplicate topic heading(s): {duplicate_topic_headings}")

        # -------------------------------------------------------------
        # 3. GENERIC STRUCTURAL TEMPLATE SEQUENCE AUDIT
        # -------------------------------------------------------------
        structural_template_families: List[str] = []
        sequence_counts: Dict[str, int] = {}
        for top_title, seq in topic_structure_map.items():
            if len(seq) >= 3:
                seq_str = " -> ".join(seq)
                sequence_counts[seq_str] = sequence_counts.get(seq_str, 0) + 1

        for seq_str, count in sequence_counts.items():
            if count >= 3:
                family_msg = f"Repeated generic section structure sequence '{seq_str}' across {count} topics."
                structural_template_families.append(family_msg)
                blocking_reasons.append(family_msg)

        # -------------------------------------------------------------
        # 4. FULL PROSE REPETITION AUDIT (EVERY PROSE PARAGRAPH)
        # -------------------------------------------------------------
        evaluated_prose_count = len(prose_paragraphs)
        exact_duplicate_paragraphs = 0
        near_duplicate_paragraphs = 0

        exact_hash_locs: Dict[str, str] = {}
        stored_tokens_list: List[Dict[str, Any]] = []

        for p in prose_paragraphs:
            norm = p["normalized"]
            loc = p["location"]
            p_hash = hashlib.sha256(norm.encode("utf-8")).hexdigest()

            if p_hash in exact_hash_locs:
                exact_duplicate_paragraphs += 1
                continue
            else:
                exact_hash_locs[p_hash] = loc

            tokens = set(norm.split())
            is_near = False
            for item in stored_tokens_list:
                s_tokens = item["tokens"]
                inter = len(tokens & s_tokens)
                union = len(tokens | s_tokens)
                sim = inter / union if union > 0 else 0.0
                if sim >= self.near_threshold:
                    near_duplicate_paragraphs += 1
                    is_near = True
                    break

            if not is_near:
                stored_tokens_list.append({"tokens": tokens, "location": loc})

        denom = max(1, evaluated_prose_count)
        exact_rate = exact_duplicate_paragraphs / denom
        near_rate = near_duplicate_paragraphs / denom

        if exact_rate > 0.02 or exact_duplicate_paragraphs > 0:
            blocking_reasons.append(
                f"Prose paragraph exact duplicate audit failed: {exact_duplicate_paragraphs} exact duplicates ({exact_rate:.2%})."
            )
        if near_rate > 0.05:
            blocking_reasons.append(
                f"Prose paragraph near duplicate audit failed: {near_duplicate_paragraphs} near duplicates ({near_rate:.2%})."
            )

        # -------------------------------------------------------------
        # 5. TABLE AUDIT (DUPLICATES & MATH RENDERING ARTIFACTS)
        # -------------------------------------------------------------
        total_tables = len(doc.tables)

        has_scorecard_table = False
        if total_tables > 0:
            last_tbl = doc.tables[-1]
            last_cell = last_tbl.rows[0].cells[0].text if last_tbl.rows and last_tbl.rows[0].cells else ""
            if "scorecard" in last_cell.lower() or "metric" in last_cell.lower():
                has_scorecard_table = True

        content_tables = (total_tables - 1) if has_scorecard_table else total_tables

        exact_duplicate_tables = 0
        near_duplicate_tables = 0
        semantic_duplicate_tables = 0
        math_rendering_errors_in_tables = 0

        table_strings: List[str] = []
        table_fingerprints: Set[str] = set()

        DE_BROGLIE_TABLE_KEYWORDS = {"cricket ball", "smoke particle", "thermal neutron", "electron"}

        for t_idx, tbl in enumerate(doc.tables):
            # Exclude appendix scorecard table if last table
            if has_scorecard_table and t_idx == total_tables - 1:
                continue

            tbl_text_parts = []
            has_table_math_error = False

            for row in tbl.rows:
                row_cells = []
                for cell in row.cells:
                    c_text = cell.text.strip()
                    row_cells.append(c_text)

                    # Check for unrendered math artifacts in table cells
                    for pat in self.UNRENDERED_MATH_PATTERNS:
                        if pat.search(c_text):
                            has_table_math_error = True
                            break

                tbl_text_parts.append(" | ".join(row_cells))

            full_tbl_str = "\n".join(tbl_text_parts)
            norm_tbl = self._normalize(full_tbl_str)

            if has_table_math_error:
                math_rendering_errors_in_tables += 1

            # Exact table duplicate check
            tbl_hash = hashlib.sha256(norm_tbl.encode("utf-8")).hexdigest()
            if tbl_hash in table_fingerprints:
                exact_duplicate_tables += 1
            else:
                table_fingerprints.add(tbl_hash)

            # Semantic table check (e.g. repeated de Broglie table across topics)
            if all(k in norm_tbl for k in DE_BROGLIE_TABLE_KEYWORDS):
                if any(all(k in t_prev for k in DE_BROGLIE_TABLE_KEYWORDS) for t_prev in table_strings):
                    semantic_duplicate_tables += 1

            # Near table duplicate check
            tbl_tokens = set(norm_tbl.split())
            if tbl_tokens:
                for prev_tbl in table_strings:
                    prev_tokens = set(prev_tbl.split())
                    inter = len(tbl_tokens & prev_tokens)
                    union = len(tbl_tokens | prev_tokens)
                    sim = inter / union if union > 0 else 0.0
                    if sim >= 0.70:
                        near_duplicate_tables += 1
                        break

            table_strings.append(norm_tbl)

        if exact_duplicate_tables > 0:
            blocking_reasons.append(f"Found {exact_duplicate_tables} exact duplicate table(s) in final DOCX.")
        if semantic_duplicate_tables > 0:
            blocking_reasons.append(f"Found {semantic_duplicate_tables} repeated semantic comparison table(s) across topics.")
        if math_rendering_errors_in_tables > 0:
            blocking_reasons.append(f"Found {math_rendering_errors_in_tables} table(s) with unrendered raw math LaTeX artifacts.")

        # -------------------------------------------------------------
        # 6. FIGURE & CAPTION SEMANTIC ALIGNMENT AUDIT
        # -------------------------------------------------------------
        figure_caption_mismatches = 0
        for cap in figure_captions:
            cap_text = cap["text"]
            cap_top = cap["topic"].lower()
            cap_sub = cap["subtopic"].lower()

            # Check if caption title refers to a different topic
            top_words = [w for w in cap_top.split() if len(w) > 4]
            if cap_text and top_words:
                # If caption has explicit topic title that doesn't match current topic
                if "schematic of" in cap_text.lower():
                    caption_topic_ref = cap_text.lower().split("schematic of")[-1].strip()
                    if cap_top not in caption_topic_ref and not any(w in caption_topic_ref for w in top_words):
                        figure_caption_mismatches += 1

        if figure_caption_mismatches > 0:
            blocking_reasons.append(f"Found {figure_caption_mismatches} figure caption semantic mismatch(es).")

        # Total unrendered math artifacts count (prose + tables)
        total_math_rendering_artifacts = math_artifacts_in_text + math_rendering_errors_in_tables
        if total_math_rendering_artifacts > 0:
            blocking_reasons.append(f"Detected {total_math_rendering_artifacts} unrendered raw math LaTeX artifact(s) in document.")

        # -------------------------------------------------------------
        # 7. STRICT COUNT RECONCILIATION AUDIT (planned == assembled == rendered)
        # -------------------------------------------------------------
        rendered_manifest = CountManifest(
            chapters=len(chapter_headings),
            topics=len(topic_headings),
            sections=len(subtopic_headings),
            paragraphs=total_docx_paragraphs,
            equations=total_equations,
            tables=content_tables,
            figures=total_figures,
            words=total_words
        )

        discrepancies: List[str] = []
        count_reconciliation_valid = True

        if planned_manifest:
            # Check STRICT equality between planned and rendered
            if planned_manifest.chapters != rendered_manifest.chapters:
                discrepancies.append(f"Chapters mismatch: planned ({planned_manifest.chapters}) != rendered ({rendered_manifest.chapters})")
                count_reconciliation_valid = False
            if planned_manifest.topics != rendered_manifest.topics:
                discrepancies.append(f"Topics mismatch: planned ({planned_manifest.topics}) != rendered ({rendered_manifest.topics})")
                count_reconciliation_valid = False
            if planned_manifest.sections != rendered_manifest.sections:
                discrepancies.append(f"Sections mismatch: planned ({planned_manifest.sections}) != rendered ({rendered_manifest.sections})")
                count_reconciliation_valid = False
            if planned_manifest.tables != rendered_manifest.tables:
                discrepancies.append(f"Tables mismatch: planned ({planned_manifest.tables}) != rendered ({rendered_manifest.tables})")
                count_reconciliation_valid = False
            if planned_manifest.figures != rendered_manifest.figures:
                discrepancies.append(f"Figures mismatch: planned ({planned_manifest.figures}) != rendered ({rendered_manifest.figures})")
                count_reconciliation_valid = False

        if assembled_manifest:
            if assembled_manifest.chapters != rendered_manifest.chapters:
                discrepancies.append(f"Chapters assembled vs rendered mismatch: assembled ({assembled_manifest.chapters}) != rendered ({rendered_manifest.chapters})")
                count_reconciliation_valid = False
            if assembled_manifest.topics != rendered_manifest.topics:
                discrepancies.append(f"Topics assembled vs rendered mismatch: assembled ({assembled_manifest.topics}) != rendered ({rendered_manifest.topics})")
                count_reconciliation_valid = False
            if assembled_manifest.sections != rendered_manifest.sections:
                discrepancies.append(f"Sections assembled vs rendered mismatch: assembled ({assembled_manifest.sections}) != rendered ({rendered_manifest.sections})")
                count_reconciliation_valid = False
            if assembled_manifest.tables != rendered_manifest.tables:
                discrepancies.append(f"Tables assembled vs rendered mismatch: assembled ({assembled_manifest.tables}) != rendered ({rendered_manifest.tables})")
                count_reconciliation_valid = False
            if assembled_manifest.figures != rendered_manifest.figures:
                discrepancies.append(f"Figures assembled vs rendered mismatch: assembled ({assembled_manifest.figures}) != rendered ({rendered_manifest.figures})")
                count_reconciliation_valid = False

        if not count_reconciliation_valid:
            blocking_reasons.extend(discrepancies)

        # -------------------------------------------------------------
        # 8. FINAL PUBLICATION READINESS DECISION
        # -------------------------------------------------------------
        publication_ready = (len(blocking_reasons) == 0)

        summary_counts = rendered_manifest.to_dict()
        reconciliation_details = {
            "is_valid": count_reconciliation_valid,
            "discrepancies": discrepancies,
            "planned": planned_manifest.to_dict() if planned_manifest else None,
            "assembled": assembled_manifest.to_dict() if assembled_manifest else None,
            "rendered": rendered_manifest.to_dict()
        }

        logger.info(
            f"IndependentArtifactAuditor complete for '{docx_path}': "
            f"publication_ready={publication_ready} (blocking_issues={len(blocking_reasons)})"
        )

        return IndependentAuditResult(
            publication_ready=publication_ready,
            docx_path=docx_path,
            total_docx_paragraphs=total_docx_paragraphs,
            evaluated_prose_paragraphs=evaluated_prose_count,
            exact_duplicate_paragraphs=exact_duplicate_paragraphs,
            near_duplicate_paragraphs=near_duplicate_paragraphs,
            exact_duplicate_paragraph_rate=exact_rate,
            near_duplicate_paragraph_rate=near_rate,
            total_headings=len(headings),
            consecutive_duplicate_headings=consecutive_duplicate_headings,
            generic_headings=generic_headings_count,
            structural_template_families=structural_template_families,
            total_tables=content_tables,
            exact_duplicate_tables=exact_duplicate_tables,
            near_duplicate_tables=near_duplicate_tables,
            semantic_duplicate_tables=semantic_duplicate_tables,
            math_rendering_errors_in_tables=math_rendering_errors_in_tables,
            total_figures=total_figures,
            figure_caption_mismatches=figure_caption_mismatches,
            math_rendering_artifacts_count=total_math_rendering_artifacts,
            count_reconciliation_valid=count_reconciliation_valid,
            blocking_reasons=blocking_reasons,
            summary_counts=summary_counts,
            reconciliation_details=reconciliation_details
        )
