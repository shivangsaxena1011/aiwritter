import os
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

class DocxExporter:
    def export(self, book_title: str, toc_data: dict, content_sections: list, images: list, output_path: str) -> str:
        doc = Document()
        
        # Professional formatting: Apply to Normal style
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)

        # Title page
        title = doc.add_heading(book_title, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()

        # Table of Contents
        doc.add_heading("Table of Contents", level=1)
        for unit in toc_data.get('units', []):
            doc.add_paragraph(unit['name'], style='List Bullet')
            for topic in unit.get('topics', []):
                doc.add_paragraph("  " + topic['name'], style='List Bullet 2')
                for subtopic in topic.get('subtopics', []):
                    doc.add_paragraph("    " + subtopic, style='List Bullet 3')
        doc.add_page_break()

        # Content
        for section in content_sections:
            if section.get('is_intro'):
                doc.add_heading(section['unit'], level=1)
                doc.add_heading("Introduction", level=2)
            else:
                if section.get('is_first_in_topic'):
                    doc.add_heading(section['topic'], level=2)
                doc.add_heading(section['subtopic'], level=3)

            content = section.get('content', '')
            self._parse_markdown(doc, content)

            if section.get('image_path') and os.path.exists(section['image_path']):
                try:
                    doc.add_picture(section['image_path'], width=Inches(6.0))
                except Exception:
                    pass
            elif section.get('image_prompt'):
                p = doc.add_paragraph()
                p.add_run(f"[DIAGRAM PLACEHOLDER: {section['image_prompt']}]").italic = True

            doc.add_paragraph() # Spacing

        doc.save(output_path)
        return output_path

    def _parse_markdown(self, doc, content: str):
        lines = content.split('\n')
        in_code_block = False
        code_content = []
        
        for line in lines:
            if line.startswith('```'):
                if in_code_block:
                    p = doc.add_paragraph('\n'.join(code_content))
                    p.style = 'Normal'
                    code_content = []
                    in_code_block = False
                else:
                    in_code_block = True
                continue
            
            if in_code_block:
                code_content.append(line)
                continue
            
            line_stripped = line.strip()
            
            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('#### '):
                doc.add_heading(line[5:], level=4)
            elif line.startswith('- '):
                p = doc.add_paragraph(style='List Bullet')
                self._parse_inline_formatting(p, line[2:])
            elif re.match(r'^\d+\.\s', line_stripped):
                text = re.sub(r'^\d+\.\s', '', line_stripped)
                p = doc.add_paragraph(style='List Number')
                self._parse_inline_formatting(p, text)
            elif line_stripped.startswith('|') and line_stripped.endswith('|'):
                p = doc.add_paragraph()
                self._parse_inline_formatting(p, line_stripped)
            elif line_stripped == '':
                continue
            else:
                p = doc.add_paragraph()
                self._parse_inline_formatting(p, line_stripped)
                
    def _parse_inline_formatting(self, paragraph, text: str):
        # Extremely basic markdown inline parser
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                paragraph.add_run(part[2:-2]).bold = True
            elif part.startswith('*') and part.endswith('*'):
                paragraph.add_run(part[1:-1]).italic = True
            elif part.startswith('`') and part.endswith('`'):
                paragraph.add_run(part[1:-1])
            else:
                paragraph.add_run(part)
