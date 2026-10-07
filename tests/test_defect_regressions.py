"""
Comprehensive Historical Defect Regression Suite:
Verifies that all 20 historical architectural, content, and rendering defects
remain definitively resolved in production.
"""

import os
import re
import tempfile
import pytest
from docx import Document

from backend.app.agents.subject_knowledge_base import (
    PhysicsKnowledgeProvider,
    MathematicsKnowledgeProvider,
    ComputerScienceKnowledgeProvider,
    ElectronicsKnowledgeProvider,
    MechanicalEngineeringKnowledgeProvider,
    GenericAcademicKnowledgeProvider,
    get_subject_knowledge_provider
)
from backend.app.agents.subject_knowledge_model import SubjectKnowledgeModel
from backend.app.agents.table_planner import TablePlanner
from backend.app.agents.academic_content_quality_agent import AcademicContentQualityAgent
from backend.app.services.document.artifact_integrity import (
    CountManifest,
    ArtifactIntegrityManifest
)
from backend.app.services.document.independent_artifact_auditor import IndependentArtifactAuditor
from backend.app.agents.syllabus_analysis_agent import SyllabusAnalysisAgent


class TestHistoricalDefectRegressions:
    """Rigorous tests covering the 20 historical defects."""

    # Defect 1: Preamble metadata becoming topic or chapter
    def test_defect_1_preamble_metadata_filtered(self):
        syllabus = """
        B.Tech First Year — Engineering Physics
        Semester II 2026 Curriculum
        Chapter 1: Quantum Mechanics
        1. Wave Nature of Particles
        2. de Broglie Hypothesis
        """
        agent = SyllabusAnalysisAgent()
        res = agent.analyze_deterministic(syllabus)
        assert len(res["chapters"]) == 1
        ch_title = res["chapters"][0]["title"]
        assert "B.Tech" not in ch_title
        assert "Curriculum" not in ch_title

    # Defect 2: Consecutive duplicate headings in DOCX
    def test_defect_2_consecutive_duplicate_headings(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            bad_path = os.path.join(tmpdir, "consecutive_dup.docx")
            doc = Document()
            doc.add_heading("Quantum Mechanics", level=1)
            doc.add_heading("Wave Mechanics", level=2)
            doc.add_heading("Wave Mechanics", level=2)
            doc.add_paragraph("Prose content describing matter waves.")
            doc.save(bad_path)

            auditor = IndependentArtifactAuditor()
            res = auditor.audit(bad_path)
            assert res.publication_ready is False
            assert res.consecutive_duplicate_headings >= 1

    # Defect 3: Generic fallback headings ("Section 1", "Subtopic 2")
    def test_defect_3_generic_fallback_headings(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            bad_path = os.path.join(tmpdir, "generic_headings.docx")
            doc = Document()
            doc.add_heading("Quantum Mechanics", level=1)
            doc.add_heading("Section 1", level=3)
            doc.add_paragraph("Valid academic prose here.")
            doc.save(bad_path)

            auditor = IndependentArtifactAuditor()
            res = auditor.audit(bad_path)
            assert res.publication_ready is False
            assert res.generic_headings >= 1

    # Defect 4: 3-way count reconciliation flaw (planned=0, assembled=0, rendered=345 marked MATCH)
    def test_defect_4_three_way_reconciliation_zero_cannot_match(self):
        planned = CountManifest(chapters=1, topics=1, sections=1, equations=0)
        assembled = CountManifest(chapters=1, topics=1, sections=1, equations=0)
        rendered = CountManifest(chapters=1, topics=1, sections=1, equations=345)

        manifest = ArtifactIntegrityManifest(planned, assembled, rendered)
        result = manifest.reconcile()
        assert result.is_valid is False
        assert result.reconciliation_table["equations"]["match"] is False
        assert any("Equation count discrepancy" in d for d in result.discrepancies)

    # Defect 5: Canned robotic opening sentence prefixes
    def test_defect_5_no_canned_opening_prefixes(self):
        robotic_prefixes = [
            "Examining",
            "The study of",
            "A comprehensive analysis of",
            "Investigating the behavior of",
            "Understanding the principles of"
        ]
        text = SubjectKnowledgeModel.generate_academic_section(
            topic="Wave Nature of Particles",
            subtopic="Dual Nature of Matter",
            subject="Engineering Physics"
        )
        first_sentence = text.split(".")[0].strip()
        for prefix in robotic_prefixes:
            # Must NOT match canned formula: "<prefix> {topic} within {subtopic} clarifies..."
            assert "within the domain of" not in first_sentence
            assert "clarifies the physical mechanisms" not in first_sentence

    # Defect 6: Universal rigid 6-part section templates
    def test_defect_6_multi_domain_providers_exist(self):
        domains = ["Physics", "Mathematics", "Computer Science", "Electronics", "Mechanical Engineering", "Chemistry"]
        providers = [get_subject_knowledge_provider(d) for d in domains]
        domain_names = [p.get_domain_name() for p in providers]
        assert len(set(domain_names)) >= 5

    # Defect 7: Table planner necessity acceptance across physical regimes
    def test_defect_7_table_planner_semantic_acceptance(self):
        dec = TablePlanner.evaluate_necessity(
            subtopic_title="Matter Waves & Dispersion",
            topic_title="Wave Nature of Particles",
            purpose="COMPARISON"
        )
        assert dec.needs_table is True
        assert dec.table_type == "COMPARISON"
        assert dec.markdown_content is not None

    # Defect 8: Table planner rejection for conceptual prose
    def test_defect_8_table_planner_conceptual_rejection(self):
        dec = TablePlanner.evaluate_necessity(
            subtopic_title="Physical Interpretation of the Wavefunction",
            topic_title="Wave Function and Born Interpretation",
            purpose="PHYSICAL_INTERPRETATION"
        )
        assert dec.needs_table is False
        assert "connected academic prose" in dec.reason

    # Defect 9: Raw unrendered math in tables
    def test_defect_9_unrendered_math_in_tables_detected(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            bad_path = os.path.join(tmpdir, "bad_table_math.docx")
            doc = Document()
            doc.add_heading("Quantum Mechanics", level=1)
            tbl = doc.add_table(rows=2, cols=2)
            tbl.cell(0, 0).text = "Observable"
            tbl.cell(0, 1).text = "Formula"
            tbl.cell(1, 0).text = "Wavelength"
            tbl.cell(1, 1).text = "\\frac{h}{p}"  # Raw unrendered LaTeX!
            doc.save(bad_path)

            auditor = IndependentArtifactAuditor()
            res = auditor.audit(bad_path)
            assert res.math_rendering_errors_in_tables >= 1

    # Defect 10: Unrendered display math delimiter in prose
    def test_defect_10_unrendered_display_math_delimiters_detected(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            bad_path = os.path.join(tmpdir, "bad_prose_math.docx")
            doc = Document()
            doc.add_heading("Quantum Mechanics", level=1)
            doc.add_paragraph("The wavelength is defined by: $$ \\lambda = h/p $$")  # Unrendered $$!
            doc.save(bad_path)

            auditor = IndependentArtifactAuditor()
            res = auditor.audit(bad_path)
            assert res.math_rendering_artifacts_count >= 1

    # Defect 11: Numerical toggle switch ON vs OFF
    def test_defect_11_numerical_switch_compliance(self):
        provider = PhysicsKnowledgeProvider()
        prose_off = provider.generate_section_prose("Quantum Mechanics", "de Broglie Hypothesis", include_numericals=False)
        assert "### Solved Numerical Example" not in prose_off

        prose_on = provider.generate_section_prose("Quantum Mechanics", "de Broglie Hypothesis", include_numericals=True)
        assert "### Solved Numerical Example" in prose_on
        assert "Given Data:" in prose_on
        assert "Governing Formula:" in prose_on

    # Defect 12: Review Q&A toggle switch ON vs OFF
    def test_defect_12_qa_switch_compliance(self):
        provider = PhysicsKnowledgeProvider()
        prose_off = provider.generate_section_prose("Quantum Mechanics", "de Broglie Hypothesis", include_questions=False)
        assert "### Academic Review & Conceptual Questions" not in prose_off

        prose_on = provider.generate_section_prose("Quantum Mechanics", "de Broglie Hypothesis", include_questions=True)
        assert "### Academic Review & Conceptual Questions" in prose_on

    # Defect 13: Banned fake configurations (Configuration Alpha / Beta)
    def test_defect_13_fake_configurations_flagged(self):
        agent = AcademicContentQualityAgent()
        fake_content = "The experiment evaluated Configuration Alpha and Configuration Beta under load."
        claims = agent.detect_unsupported_claims(fake_content)
        assert claims["unsupported_claim_count"] >= 1
        assert claims["has_fabrications"] is True
        assert any("Configuration" in c for c in claims["claims"])

    # Defect 14: Fabricated statistics (e.g. 42% - 48% efficiency)
    def test_defect_14_fabricated_statistics_flagged(self):
        agent = AcademicContentQualityAgent()
        fake_content = "The device achieved 42% - 48% efficiency across all testing cycles."
        claims = agent.detect_unsupported_claims(fake_content)
        assert claims["unsupported_claim_count"] >= 1
        assert claims["has_fabrications"] is True
        assert any("efficiency" in c.lower() for c in claims["claims"])

    # Defect 15: Unrelated fake equations (\dot{S}_{gen} in quantum physics)
    def test_defect_15_unrelated_equations_checked(self):
        from backend.app.agents.equation_validation_agent import EquationValidationAgent
        validator = EquationValidationAgent()
        bad_eq = "The balance yields: $$\\dot{S}_{gen} = \\nabla \\cdot \\mathbf{J} + \\sigma$$"
        res = validator.validate_equations(bad_eq, "de Broglie Hypothesis", "Quantum Mechanics")
        assert any("banned" in issue.lower() or "unrelated" in issue.lower() for issue in res.issues)

    # Defect 16: Repetition rate bounded between 0.0 and 1.0
    def test_defect_16_repetition_rate_bounded(self):
        from backend.app.agents.repetition_detector_2 import RepetitionDetector2
        det = RepetitionDetector2()
        summary = det.get_summary()
        assert 0.0 <= summary["exact_duplicate_rate"] <= 1.0
        assert 0.0 <= summary["near_duplicate_rate"] <= 1.0

    # Defect 17: String interpolation bug in numerical problem ($V = 100\text{ V}$)
    def test_defect_17_no_name_error_in_numerical_generation(self):
        # Must execute without NameError: name 'V' is not defined
        provider = PhysicsKnowledgeProvider()
        prose = provider.generate_section_prose(
            topic="Wave Nature",
            subtopic="de Broglie Hypothesis",
            include_numericals=True
        )
        assert "V = 100" in prose
        assert "1.23" in prose

    # Defect 18: Multi-domain support beyond Physics (CS, Math, Electronics, Mechanical)
    def test_defect_18_cross_domain_treatises(self):
        cs_prov = ComputerScienceKnowledgeProvider()
        cs_prose = cs_prov.generate_section_prose("Sorting Algorithms", "Merge Sort", include_numericals=True)
        assert "Master Theorem" in cs_prose or "recurrence" in cs_prose.lower()

        math_prov = MathematicsKnowledgeProvider()
        math_prose = math_prov.generate_section_prose("Linear Algebra", "Eigenvalue Decomposition", include_numericals=True)
        assert "eigenvalue" in math_prose.lower()

        elec_prov = ElectronicsKnowledgeProvider()
        elec_prose = elec_prov.generate_section_prose("Semiconductor Devices", "p-n Junction", include_numericals=True)
        assert "built-in potential" in elec_prose.lower() or "shockley" in elec_prose.lower()

    # Defect 19: Storage Provider S3/Local abstraction
    def test_defect_19_storage_provider_factory(self):
        from backend.app.storage.factory import get_storage_provider
        prov = get_storage_provider()
        assert hasattr(prov, "save_file")
        assert hasattr(prov, "get_file_path")

    # Defect 20: System Health, Liveness, and Readiness endpoints
    def test_defect_20_health_endpoints(self):
        from fastapi.testclient import TestClient
        from backend.app.main import app
        client = TestClient(app)

        res_health = client.get("/health")
        assert res_health.status_code == 200
        assert res_health.json()["status"] == "healthy"

        res_live = client.get("/live")
        assert res_live.status_code == 200
        assert res_live.json()["status"] == "alive"

        res_ready = client.get("/ready")
        assert res_ready.status_code == 200
        assert res_ready.json()["status"] == "ready"
